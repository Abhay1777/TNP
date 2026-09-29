"""Export Placement Opportunities and Company Master data in canonical 36-column format."""

import csv
import io
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from placements.bulk_import.constants import (
    CANONICAL_COLUMNS,
    DEPARTMENT_COLUMNS,
)
from placements.models import PlacementOpportunity


def build_row_data(index: int, opp: PlacementOpportunity) -> dict:
    """Transform a PlacementOpportunity model instance into a canonical 36-column dictionary."""
    company = opp.company
    departments = opp.eligible_departments or []

    # Job profiles (up to 5)
    profiles = opp.job_profiles or []
    profile_dict = {}
    for i in range(1, 6):
        profile_dict[f"Job Profile {i}"] = profiles[i - 1] if len(profiles) >= i else ""

    # Skills (up to 8)
    skills = opp.skills or []
    skill_dict = {}
    for i in range(1, 9):
        skill_dict[f"Skillset Required {i}"] = skills[i - 1] if len(skills) >= i else ""

    # Department flags
    dept_dict = {}
    for dept in DEPARTMENT_COLUMNS:
        dept_dict[dept] = "Yes" if dept in departments else "No"

    # Emolument display
    emolument = opp.emolument_raw
    if not emolument and opp.emolument_value:
        emolument = f"{opp.emolument_value} {opp.emolument_unit}"

    row = {
        "Sr. No.": index,
        "Batch": opp.batch,
        "Name of the Company": company.name if company else "",
        "Eligibility Criteria": opp.eligibility_criteria or "",
        **dept_dict,
        "Designation": opp.designation or "",
        "Tech / Non-Tech": opp.tech_nontech or "",
        **profile_dict,
        **skill_dict,
        "Emolument (CTC)": emolument or "",
        "Selection Process": opp.selection_process or "",
        "Company Website": (company.website if company else ""),
        "Placement / Internship": opp.placement_internship or "",
        "Eligible Department": ", ".join(departments),
        "No Of Offers": opp.number_of_offers or "",
    }
    return row


def export_opportunities_to_csv(queryset) -> io.StringIO:
    """Generate CSV string buffer containing opportunities in canonical 36-column format."""
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=CANONICAL_COLUMNS)
    writer.writeheader()

    # Apply select_related; only reorder if the queryset has not already been sliced
    qs = queryset.select_related("company")
    if not qs.query.is_sliced:
        qs = qs.order_by("-batch", "company__name")

    for idx, opp in enumerate(qs, start=1):
        writer.writerow(build_row_data(idx, opp))

    output.seek(0)
    return output


def export_opportunities_to_excel(queryset) -> io.BytesIO:
    """Generate professional Excel workbook with TCET header styling and 36 canonical columns."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Placement Records"

    # Styling definitions
    header_fill = PatternFill(start_color="153F74", end_color="153F74", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    data_font = Font(name="Arial", size=9)
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")

    # Write Header Row
    ws.append(CANONICAL_COLUMNS)
    for col_num in range(1, len(CANONICAL_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
    ws.row_dimensions[1].height = 28

    # Write Data Rows
    # Apply select_related; only reorder if the queryset has not already been sliced
    qs = queryset.select_related("company")
    if not qs.query.is_sliced:
        qs = qs.order_by("-batch", "company__name")

    for idx, opp in enumerate(qs, start=1):
        row_dict = build_row_data(idx, opp)
        row_values = [row_dict.get(col, "") for col in CANONICAL_COLUMNS]
        ws.append(row_values)

        row_num = ws.max_row
        ws.row_dimensions[row_num].height = 20

        # Alternate row fill
        row_fill = PatternFill(
            start_color="F9FAFB" if idx % 2 == 0 else "FFFFFF",
            end_color="F9FAFB" if idx % 2 == 0 else "FFFFFF",
            fill_type="solid",
        )

        for col_num in range(1, len(CANONICAL_COLUMNS) + 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.font = data_font
            cell.fill = row_fill
            cell.border = thin_border

            # Alignment logic
            col_name = CANONICAL_COLUMNS[col_num - 1]
            if col_name in ["Sr. No.", "Batch", "Placement / Internship", "No Of Offers"] or col_name in DEPARTMENT_COLUMNS:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 38)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
